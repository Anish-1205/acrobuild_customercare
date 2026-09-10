from __future__ import annotations

import json
import os
import re
import time
import uuid
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any

import requests
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# The 200 prompts are committed here (id, category, prompt, expected). The old
# run extracted them from a local PDF; that path no longer exists on every
# machine, so the prompt set is version-controlled instead.
_ROOT = Path(__file__).resolve().parent.parent
PROMPTS_PATH = Path(os.getenv("EVAL_PROMPTS_PATH", _ROOT / "scripts" / "eval_prompts.json"))
OUTPUT_PATH = Path(os.getenv("EVAL_OUTPUT_PATH", "customer_support_chatbot_200_prompt_evaluation.xlsx"))
CACHE_PATH = Path(os.getenv("EVAL_CACHE_PATH", "data/chatbot_200_prompt_results.json"))
API_URL = os.getenv("EVAL_API_URL", "http://127.0.0.1:8000/api/support/assist")
MAX_WORKERS = int(os.getenv("EVAL_MAX_WORKERS", "4"))
TIMEOUT_SECONDS = int(os.getenv("EVAL_TIMEOUT_SECONDS", "60"))


def clean(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def extract_cases() -> list[dict[str, Any]]:
    raw = json.loads(PROMPTS_PATH.read_text(encoding="utf-8"))
    cases = [
        {
            "id": int(item["id"]),
            "category": clean(item["category"]),
            "prompt": clean(item["prompt"]),
            "expected": clean(item.get("expected")),
        }
        for item in raw
    ]
    cases.sort(key=lambda item: item["id"])
    ids = [item["id"] for item in cases]
    if ids != list(range(1, len(cases) + 1)):
        raise RuntimeError(f"{PROMPTS_PATH} must hold contiguous IDs from 1; got {ids}")
    return cases


def send_prompt(issue: str, conversation_id: str, history: list[dict[str, str]] | None = None) -> dict[str, Any]:
    started = time.perf_counter()
    response = requests.post(
        API_URL,
        json={
            "issue": issue,
            "conversation_id": conversation_id,
            "conversation_messages": history or [],
            "prefer_fast_response": True,
            "prefer_qwen_response": True,
            "limit": 2,
        },
        timeout=TIMEOUT_SECONDS,
    )
    elapsed_ms = round((time.perf_counter() - started) * 1000)
    response.raise_for_status()
    payload = response.json()
    payload["latency_ms"] = elapsed_ms
    return payload


def run_case(case: dict[str, Any]) -> dict[str, Any]:
    prompt = case["prompt"]
    conversation_id = f"eval-{case['id']}-{uuid.uuid4().hex[:10]}"
    previous_answer = ""
    tested_prompt = prompt
    follow_up = re.match(r"Previous:\s*(.*?)\s*Current:\s*(.*)", prompt, flags=re.IGNORECASE)
    if follow_up:
        previous_prompt = clean(follow_up.group(1))
        tested_prompt = clean(follow_up.group(2))
        previous_payload = send_prompt(previous_prompt, conversation_id)
        previous_answer = clean(previous_payload.get("answer"))
    payload = send_prompt(tested_prompt, conversation_id)
    evaluation = payload.get("rag_evaluation") or {}
    metrics = evaluation.get("metrics") or {}
    return {
        **case,
        "tested_prompt": tested_prompt,
        "previous_answer": previous_answer,
        "answer": clean(payload.get("answer")),
        "agent_mode": clean(payload.get("agent_mode")),
        "source_label": clean(payload.get("source_label")),
        "source_status": clean(payload.get("source_status")),
        "latency_ms": payload.get("latency_ms", 0),
        "retrieval_mode": clean(payload.get("retrieval_mode")),
        "api_call_count": len(payload.get("data_api_calls") or []),
        "overall_quality": (metrics.get("overall_quality") or {}).get("percent"),
        "answer_relevance": (metrics.get("answer_relevance") or {}).get("percent"),
        "groundedness": (metrics.get("groundedness") or {}).get("percent"),
        "context_precision": (metrics.get("context_precision") or {}).get("percent"),
        "retrieval_hit_rate": (metrics.get("retrieval_hit_rate") or {}).get("percent"),
        "error": "",
    }


def contains_any(text: str, terms: tuple[str, ...]) -> bool:
    return any(term in text for term in terms)


def judge(result: dict[str, Any]) -> tuple[str, str]:
    category = result["category"]
    prompt = result["tested_prompt"].lower()
    answer = result["answer"].lower()
    if result.get("error"):
        return "NOT WORKING", f"Request failed: {result['error']}"
    if not answer:
        return "NOT WORKING", "No answer returned."
    unrelated = contains_any(answer, (
        "business-hours and after-hours expectations",
        "need more help and send us your email",
        "the company listed in the latest system is gbk",
    ))
    if unrelated:
        return "NOT WORKING", "Response is unrelated to the requested intent."

    if category == "UNKNOWN QUERY":
        safety_terms = (
            "cannot", "can't", "do not", "don't", "unable", "outside", "not connected",
            "not available", "no access", "private", "qualified", "professional", "emergency",
            "rephrase", "unsupported", "not able", "won't", "will not", "assist with",
        )
        if contains_any(answer, safety_terms):
            return "WORKING", "Safely declines, redirects, or asks for clarification."
        return "NOT WORKING", "Answered an unsupported/current/sensitive request without a clear limitation."

    if category == "GENERAL CHAT":
        if result["id"] == 18 and not re.search(r"\b200\b", answer):
            return "NOT WORKING", "Incorrect or missing calculation; expected 200."
        if result["id"] == 17 and not re.search(r"[\u0900-\u097f]", result["answer"]):
            return "NOT WORKING", "Hindi translation was not produced."
        if result["id"] == 20:
            numbered = len(re.findall(r"(?:^|\s)[1-3][.)]", result["answer"]))
            bullets = result["answer"].count("•") + result["answer"].count("- ")
            if max(numbered, bullets) < 3:
                return "NOT WORKING", "Did not clearly provide the requested three tips."
        return "WORKING", "Relevant conversational response returned."

    if "[project" in prompt or "[company name]" in prompt or "[developer name]" in prompt:
        placeholder_safe = contains_any(answer, (
            "project name", "company name", "developer name", "which project", "please share",
            "not available", "unavailable", "could not", "cannot", "tell me",
        ))
        if placeholder_safe:
            return "WORKING", "Safely requests a real entity or reports unavailable data instead of inventing it."
        if "project a" in answer or "project b" in answer or "company name" in answer:
            return "NOT WORKING", "Treated a placeholder as a real entity or invented unsupported details."

    if category == "COMPANY INFORMATION":
        relevant = contains_any(answer, (
            "acrobuild", "company", "project", "support", "sales", "office", "contact", "privacy",
            "terms", "complaint", "partner", "sustain", "testimonial", "career", "opening",
            "available", "unavailable", "documented", "cannot", "could not",
        ))
        return ("WORKING", "Company-related answer or transparent limitation returned.") if relevant else ("NOT WORKING", "Answer does not address the company-information request.")

    if category == "PROPERTY SEARCH":
        relevant = contains_any(answer, (
            "project", "home", "flat", "apartment", "villa", "plot", "bhk", "property", "location",
            "budget", "available", "no matching", "please share", "which city", "tell me",
        ))
        return ("WORKING", "Search result, no-match result, or useful filter clarification returned.") if relevant else ("NOT WORKING", "No relevant property-search result or clarification.")

    if category == "PROPERTY DETAILS":
        relevant = contains_any(answer, (
            "project", "property", "location", "configuration", "bhk", "size", "area", "amenit",
            "possession", "developer", "registration", "tower", "unit", "floor plan", "parking",
            "maintenance", "security", "school", "hospital", "metro", "construction", "payment",
            "document", "brochure", "please share", "which project", "not available", "unavailable",
        ))
        return ("WORKING", "Project detail or transparent request for missing project data returned.") if relevant else ("NOT WORKING", "Answer does not address the requested property detail.")

    if category == "PROPERTY COMPARISON":
        relevant = contains_any(answer, (
            "compare", "comparison", "project", "both", "cheaper", "price", "location", "amenit",
            "size", "possession", "metro", "maintenance", "parking", "family", "senior", "investment",
            "approval", "pros", "cons", "budget", "priority", "please share", "not available",
        ))
        return ("WORKING", "Comparison, conditional recommendation, or needed clarification returned.") if relevant else ("NOT WORKING", "No relevant comparison or clarification was returned.")

    if category == "LIVE AVAILABILITY":
        relevant = contains_any(answer, (
            "available", "availability", "inventory", "unit", "sold", "flat", "home", "villa",
            "project", "floor", "site visit", "hold", "waiting list", "updated", "refresh",
            "please share", "not available", "cannot",
        ))
        return ("WORKING", "Live-availability result or transparent workflow limitation returned.") if relevant else ("NOT WORKING", "Answer does not address live availability.")

    if category == "LIVE PRICING":
        relevant = contains_any(answer, (
            "price", "pricing", "inr", "₹", "lakh", "crore", "cost", "rate", "quote", "tax",
            "registration", "parking", "charge", "deposit", "discount", "down payment", "emi",
            "interest", "exchange", "guarantee", "valid", "project", "unit", "please share",
            "not available", "cannot",
        ))
        return ("WORKING", "Pricing/calculation response or transparent request for required inputs returned.") if relevant else ("NOT WORKING", "Answer does not address pricing or the required calculation.")

    if category == "FOLLOW UP QUERY":
        if not result.get("previous_answer"):
            return "NOT WORKING", "Previous turn did not produce an answer for context testing."
        relevant_terms = {
            168: ("cheap", "affordable", "price"), 169: ("amenit",), 170: ("metro", "closer"),
            171: ("percent", "15", "lakh"), 172: ("largest", "size", "unit"), 173: ("3bhk", "3 bhk", "second"),
            174: ("confirm", "date", "possession"), 175: ("kondapur",), 176: ("1.1 crore", "110", "budget"),
            177: ("under-construction", "under construction"), 178: ("compare", "project b"),
            179: ("hospital",), 180: ("ह", "Hindi"), 181: ("short",), 182: ("parking",),
            183: ("east", "facing"), 184: ("site visit", "saturday", "book"), 185: ("expensive", "option"),
            186: ("priority", "choose", "cheaper", "closer"), 187: ("villa",),
        }
        terms = tuple(str(term).lower() for term in relevant_terms.get(result["id"], ()))
        if terms and not contains_any(answer, terms):
            return "NOT WORKING", "Follow-up answer did not preserve or apply the requested context."
        return "WORKING", "Follow-up response addresses the new turn using conversation context."

    return "NOT WORKING", "Could not verify the expected behavior automatically."


def load_cache() -> dict[str, dict[str, Any]]:
    if not CACHE_PATH.exists():
        return {}
    try:
        return {
            str(item["id"]): item
            for item in json.loads(CACHE_PATH.read_text(encoding="utf-8"))
            if not clean(item.get("error"))
        }
    except Exception:
        return {}


def save_cache(results: list[dict[str, Any]]) -> None:
    CACHE_PATH.parent.mkdir(parents=True, exist_ok=True)
    CACHE_PATH.write_text(json.dumps(results, ensure_ascii=False, indent=2), encoding="utf-8")


def create_workbook(results: list[dict[str, Any]]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    total = len(results)
    working = sum(item["status"] == "WORKING" for item in results)
    not_working = total - working
    summary.append(["Customer Support Chatbot - 200 Prompt Evaluation"])
    summary.append([])
    summary.append(["Metric", "Value"])
    summary.append(["Total prompts", total])
    summary.append(["Working", working])
    summary.append(["Not working", not_working])
    summary.append(["Pass rate", working / total if total else 0])
    summary.append(["Average latency (ms)", round(sum(item.get("latency_ms", 0) for item in results) / total) if total else 0])
    summary["A1"].font = Font(size=16, bold=True)
    summary["A3"].font = summary["B3"].font = Font(bold=True, color="FFFFFF")
    summary["A3"].fill = summary["B3"].fill = PatternFill("solid", fgColor="173F35")
    summary["B7"].number_format = "0.0%"
    summary.column_dimensions["A"].width = 32
    summary.column_dimensions["B"].width = 20

    summary.append([])
    summary.append(["Category", "Total", "Working", "Not working", "Pass rate"])
    category_header = summary.max_row
    for cell in summary[category_header]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="173F35")
    categories = []
    for item in results:
        if item["category"] not in categories:
            categories.append(item["category"])
    for category in categories:
        rows = [item for item in results if item["category"] == category]
        passed = sum(item["status"] == "WORKING" for item in rows)
        summary.append([category, len(rows), passed, len(rows) - passed, passed / len(rows)])
        summary.cell(summary.max_row, 5).number_format = "0.0%"

    details = wb.create_sheet("Prompt Results")
    headers = [
        "ID", "Category", "Original Test Prompt", "Prompt Sent", "Expected POC Behavior",
        "Chatbot Answer", "Status", "Evaluation Reason", "Latency (ms)", "Agent Mode",
        "Source", "Source Status", "Retrieval Mode", "API Calls", "Overall Quality %",
        "Answer Relevance %", "Groundedness %", "Context Precision %", "Retrieval Hit Rate %",
        "Previous Turn Answer"
    ]
    details.append(headers)
    for item in results:
        details.append([
            item["id"], item["category"], item["prompt"], item["tested_prompt"], item["expected"],
            item["answer"], item["status"], item["reason"], item.get("latency_ms", 0),
            item.get("agent_mode", ""), item.get("source_label", ""), item.get("source_status", ""),
            item.get("retrieval_mode", ""), item.get("api_call_count", 0), item.get("overall_quality"),
            item.get("answer_relevance"), item.get("groundedness"), item.get("context_precision"),
            item.get("retrieval_hit_rate"), item.get("previous_answer", "")
        ])
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    header_fill = PatternFill("solid", fgColor="173F35")
    pass_fill = PatternFill("solid", fgColor="D9EAD3")
    fail_fill = PatternFill("solid", fgColor="F4CCCC")
    for cell in details[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [7, 24, 48, 42, 48, 70, 16, 48, 14, 20, 28, 18, 18, 11, 18, 20, 18, 20, 20, 60]
    for index, width in enumerate(widths, 1):
        details.column_dimensions[get_column_letter(index)].width = width
    for row in details.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[6].fill = pass_fill if row[6].value == "WORKING" else fail_fill
    for col in range(15, 20):
        details.conditional_formatting.add(
            f"{get_column_letter(col)}2:{get_column_letter(col)}{details.max_row}",
            ColorScaleRule(start_type="min", start_color="F8696B", mid_type="percentile", mid_value=50,
                           mid_color="FFEB84", end_type="max", end_color="63BE7B")
        )

    notes = wb.create_sheet("Methodology")
    methodology = [
        ["Evaluation methodology"],
        ["Execution", "Prompts were sent to the live local /api/support/assist endpoint with the same Qwen-enabled setting used by the widget."],
        ["Follow-up prompts", "Items 168-187 were executed as two-turn conversations using one conversation ID."],
        ["Placeholders", "[Project A], [Project B], [Company Name], and [Developer Name] were kept as written. Asking for a real name or reporting unavailable data counts as safe behavior; invented details fail."],
        ["Verdict", "WORKING means the response addressed the expected intent, safely requested missing information, or transparently reported a limitation. NOT WORKING means it was irrelevant, unsafe, fabricated, empty, or missed a required operation."],
        ["RAG metrics", "Online proxy metrics are copied from the chatbot response. Recall, MRR, and NDCG are omitted because the test PDF does not provide labelled expected source documents."],
        ["Important", "Automated verdicts should be reviewed before production sign-off, especially subjective comparison and natural-language formatting cases."],
    ]
    for row in methodology:
        notes.append(row)
    notes["A1"].font = Font(size=15, bold=True)
    notes.column_dimensions["A"].width = 24
    notes.column_dimensions["B"].width = 115
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(OUTPUT_PATH)


def main() -> None:
    cases = extract_cases()
    cache = load_cache()
    results: list[dict[str, Any]] = list(cache.values())
    pending = [case for case in cases if str(case["id"]) not in cache]
    print(f"Extracted {len(cases)} cases; {len(pending)} pending")
    if pending:
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = {executor.submit(run_case, case): case for case in pending}
            for index, future in enumerate(as_completed(futures), 1):
                case = futures[future]
                try:
                    result = future.result()
                except Exception as exc:
                    result = {
                        **case, "tested_prompt": case["prompt"], "previous_answer": "", "answer": "",
                        "agent_mode": "", "source_label": "", "source_status": "", "latency_ms": 0,
                        "retrieval_mode": "", "api_call_count": 0, "overall_quality": None,
                        "answer_relevance": None, "groundedness": None, "context_precision": None,
                        "retrieval_hit_rate": None, "error": f"{type(exc).__name__}: {exc}",
                    }
                results = [item for item in results if item["id"] != case["id"]] + [result]
                results.sort(key=lambda item: item["id"])
                save_cache(results)
                print(f"[{index}/{len(pending)}] ID {case['id']} completed", flush=True)
    for result in results:
        result["status"], result["reason"] = judge(result)
    results.sort(key=lambda item: item["id"])
    save_cache(results)
    create_workbook(results)
    print(f"Saved {OUTPUT_PATH.resolve()}")


if __name__ == "__main__":
    main()
